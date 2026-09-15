from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUTPUT = r"D:\project\Dify 课堂互动智能体\课堂互动智能体系统说明.docx"

BLACK = "000000"
HEADER_FILL = "1F4E79"
HEADER_TEXT = "FFFFFF"
ALT_FILL = "F2F6FA"
BORDER = "D9D9D9"


def set_run_font(run, size=10.5, bold=False, color=BLACK, name="Microsoft YaHei"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def set_paragraph_format(paragraph, before=0, after=6, line=1.18, align=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        fmt.alignment = align


def add_body(document, text, bold_lead=None, after=6):
    paragraph = document.add_paragraph()
    set_paragraph_format(paragraph, after=after)
    if bold_lead and text.startswith(bold_lead):
        first = paragraph.add_run(bold_lead)
        set_run_font(first, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    set_paragraph_format(paragraph, after=3, line=1.15)
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_number(document, text):
    paragraph = document.add_paragraph(style="List Number")
    set_paragraph_format(paragraph, after=3, line=1.15)
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(document, headers, rows, widths):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Cm(widths[index])
        set_cell_shading(cell, HEADER_FILL)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_paragraph_format(paragraph, after=0, line=1.0)
        run = paragraph.add_run(header)
        set_run_font(run, size=10, bold=True, color=HEADER_TEXT)
    set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.width = Cm(widths[col_index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, ALT_FILL)
            paragraph = cell.paragraphs[0]
            set_paragraph_format(paragraph, after=0, line=1.1)
            run = paragraph.add_run(str(value))
            set_run_font(run, size=9.5)
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_mono_block(document, text):
    paragraph = document.add_paragraph()
    set_paragraph_format(paragraph, before=3, after=8, line=1.05)
    run = paragraph.add_run(text)
    set_run_font(run, size=9, name="Consolas")
    return paragraph


def set_document_defaults(document):
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.25)
    section.right_margin = Cm(2.25)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    for style_name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12.5)):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def build_document():
    document = Document()
    set_document_defaults(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_format(title, after=8, line=1.0)
    run = title.add_run("课堂互动智能体系统说明与数据输入输出")
    set_run_font(run, size=22, bold=True)

    subtitle = document.add_paragraph()
    set_paragraph_format(subtitle, after=12, line=1.15)
    run = subtitle.add_run("面向系统架构师的项目运作说明")
    set_run_font(run, size=12, bold=False)

    add_body(
        document,
        "本文档说明当前课堂互动智能体的目标、角色、完整运作流程，以及当前文件夹中已经存在的数据和每轮输入输出。当前实现使用 Markdown 和 JSON 文件承载数据，数据库、接口、存储引擎和并发方案由系统架构师重新规划。",
        after=8,
    )
    add_body(
        document,
        "阅读重点：先看系统如何运作，再看输入和输出数据，最后看架构师必须保留的规则和可以自行决定的边界。本文不要求架构师照搬当前文件布局。",
        after=10,
    )

    document.add_heading("1 系统目标与边界", level=1)
    add_body(
        document,
        "智能体面向单个学生的课堂互动场景，由主持人控制课堂节奏，学生在上课过程中标记疑问，片段结束后由 AI 针对标记点答疑，并把学习证据持续写入学生掌握档案。",
    )
    add_body(
        document,
        "当前版本是本地原型，目标是先验证“主持人 + 标记点 + 掌握档案”的运作闭环。它不负责视频播放器本身，也不负责老师端和学生 workspace 的最终界面。",
    )
    add_body(
        document,
        "现有文件系统只是当前实现的临时载体，不代表最终数据库设计。架构师可以替换为关系数据库、文档数据库、对象存储或混合方案，但需要保留本文第 8 节的架构约束。",
    )

    document.add_heading("2 系统角色", level=1)
    add_table(
        document,
        ["角色", "职责", "产生或维护的数据"],
        [
            ("老师/主持人", "控制上课开始、片段播放、片段结束和下课；维护课程目标与知识点。", "课堂事件、老师目标、课程内容、知识点目录"),
            ("学生", "上课、打标记、提问、回答问题、手动修改标签。", "学生消息、标记点、上传文件、掌握变化"),
            ("AI Agent", "主持人模式下控制节奏；答疑模式下提问、提示并判断掌握。", "AI 回复、会话状态、掌握更新"),
            ("工作流", "接收事件、读取数据、组装上下文、调用模型并写回。", "执行结果、读写调度和状态合并"),
            ("存储层", "保存课程规则、运行状态、日志和历史。", "文件、JSON、数据库记录或对象存储"),
        ],
        [2.0, 7.1, 7.2],
    )

    document.add_heading("3 完整运作流程", level=1)
    add_body(
        document,
        "下面的流程是当前产品逻辑的完整链路。每一步结束都有可观察的输出，工作流不在没有证据的情况下推进掌握状态。",
    )
    add_table(
        document,
        ["阶段", "输入", "处理", "输出"],
        [
            ("课前准备", "老师目标、课程内容、知识点目录、课程片段", "建立规则层和知识目录", "可读取的课程配置"),
            ("上课开始", "lesson_start 或手动命令", "AI 以主持人身份介绍课程", "开场回复和初始会话状态"),
            ("片段播放", "segment_start 和 active_segment_id", "记录当前片段，不进入答题", "播放中的会话状态"),
            ("课堂提问", "学生问题或标记操作", "不打断课堂，生成 class-point", "新的标记点和简短确认"),
            ("片段结束", "segment_end 和 segment_id", "总结片段，过滤本片段开放标记点", "标记点列表和选择请求"),
            ("标记点答疑", "学生选择 point_id", "读取片段与知识点，进行提问、提示和判断", "问答、标记状态、掌握更新"),
            ("课堂结束", "lesson_end 或 /结束", "汇总证据并更新长期记录", "结束摘要和归档状态"),
        ],
        [2.4, 4.5, 4.8, 4.6],
    )

    document.add_heading("4 输入数据", level=1)
    add_body(
        document,
        "输入分为老师提供的数据、学生产生的数据、课堂事件和外部材料。架构师可以自行确定接口、校验、版本和权限，但需要保证数据能追溯到来源。",
    )
    add_table(
        document,
        ["输入来源", "数据内容", "用途", "当前载体"],
        [
            ("老师", "TMISSION、SMISSION、LESSON-INTERACTION", "定义课堂目标、内容和难点", "Markdown 文件"),
            ("老师", "KNOWLEDGE-BASE、知识点 KP-xxx", "建立稳定的掌握档案坐标系", "Markdown 文件"),
            ("老师", "课程片段 seg-xxx", "把课件切成可答疑的小块", "JSON 文件"),
            ("学生", "文字消息、选择题、普通发言", "驱动主持人和答疑流程", "Chat 输入"),
            ("学生", "标记点与标签", "表达完全不懂、走神或部分理解", "JSON 文件"),
            ("学生", "上传文件和手动标签修改", "补充学习材料和修正判断", "上传文件与 JSON"),
            ("平台", "lesson_start、segment_start、segment_end、lesson_end", "控制课堂节奏", "当前用命令模拟，未来接口"),
            ("老师/平台", "PDF、网址、课件、接口材料", "作为备课和答疑依据", "接口预留"),
        ],
        [2.3, 5.1, 5.4, 3.5],
    )

    document.add_heading("5 输出数据", level=1)
    add_body(
        document,
        "输出数据分为给学生看的回复、给工作流使用的当前状态、给掌握档案使用的长期数据和给老师查看的汇总记录。",
    )
    add_table(
        document,
        ["输出", "内容", "写入方式", "生命周期"],
        [
            ("AI 回复", "主持人总结、答疑、提示、追问", "返回 Chat", "同步显示，同时进入问答记录"),
            ("会话状态", "host_phase、当前片段、当前问题、尝试次数", "覆盖当前会话", "当前课堂有效"),
            ("标记点状态", "open、已解决、延后", "更新标记记录", "长期保存"),
            ("掌握状态", "未检测、掌握、部分掌握、未掌握", "更新当前状态", "长期保存"),
            ("掌握历史", "旧状态、新状态、证据、时间", "追加", "只追加，不覆盖"),
            ("完整问答", "学生、主持人、AI 的每条消息", "追加", "长期保存"),
            ("长期记录", "NOTES、GLOSSARY、LEARNING-RECORD", "有证据时更新", "跨课程保存"),
        ],
        [2.5, 5.7, 4.1, 4.0],
    )

    document.add_heading("6 当前文件布局", level=1)
    add_body(
        document,
        "当前文件布局只是实现参考。规则、状态、标记和学生数据分开存放，目的是让写入责任和读取范围清楚。",
    )
    add_table(
        document,
        ["目录", "用途", "当前数据", "读写边界"],
        [
            ("class agent", "老师规则与稳定知识点目录", "SKILL、FORMAT、KNOWLEDGE-BASE", "AI 只读"),
            ("class-point", "课程片段和学生标记点", "segments/*.json、points/*.json", "片段只读，标记可写"),
            ("teach test", "本课运行状态", "七个 Markdown 状态文件", "按事件更新"),
            ("student-workspace/data", "学生长期数据", "掌握状态、历史、问答日志", "工作流和学生端共享"),
        ],
        [3.1, 4.9, 5.5, 2.8],
    )

    document.add_heading("7 数据关系", level=1)
    add_body(
        document,
        "数据关系围绕“老师知识点目录 -> 课程片段 -> 学生标记 -> 答疑证据 -> 掌握状态与历史”展开。",
    )
    add_mono_block(
        document,
        "老师知识点目录 (KP-xxx)\n"
        "    -> 课程片段 (seg-xxx) 引用若干知识点\n"
        "    -> 学生标记点 (point-xxx) 属于一个片段\n"
        "    -> 答疑证据产生掌握更新\n"
        "    -> 当前掌握状态 + 每次变化历史\n\n"
        "学生和 AI 的每轮消息\n"
        "    -> 追加到完整问答记录\n"
        "    -> 有证据时关联知识点并触发掌握更新",
    )
    add_body(
        document,
        "这套关系的关键点是：掌握状态挂在稳定知识点上，而不是挂在某一次对话或某一个文件上。课程片段和标记点负责提供上下文，掌握状态和历史负责跨课追踪。",
    )

    document.add_heading("8 架构师必须保留的规则", level=1)
    for text in [
        "老师规则和知识点目录由老师维护，AI 只读。",
        "完整问答只追加，不覆盖，不把聊天记录当成掌握状态。",
        "掌握状态和历史分开保存：当前状态用于查询，历史用于追踪变化。",
        "掌握状态由答疑证据、标记映射或学生手动修改触发。",
        "学生自报“掌握/没掌握”不等同于 AI 确认的掌握证据。",
        "每个标记点必须有状态、来源片段和回到原始位置的信息。",
        "每个知识点必须有稳定 ID，掌握历史依赖它进行长期关联。",
        "课程片段只加载相关知识点，避免把整门课内容一起读取。",
        "当前会话状态和长期学生档案分开，不能混成一个日志文件。",
        "平台事件、来源链接和老师材料接口要能替换当前样例数据。",
    ]:
        add_bullet(document, text)

    document.add_heading("9 架构师可以自行决定的部分", level=1)
    add_table(
        document,
        ["决策项", "需要回答的问题"],
        [
            ("存储模型", "关系数据库、文档数据库还是混合存储；文件是否继续保留。"),
            ("ID 设计", "业务 ID、UUID、组合键如何选择，是否需要多租户隔离。"),
            ("规则存储", "规则和知识点目录继续放文件，还是入库并做版本管理。"),
            ("状态边界", "哪些数据属于工作流运行态，哪些属于学生长期档案。"),
            ("接口协议", "老师端、学生端、视频平台和工作流之间用同步接口还是事件。"),
            ("掌握历史", "如何归档、查询、分区以及处理长期增长。"),
            ("多学生并发", "会话、标记、掌握更新如何隔离和加锁。"),
            ("学生 workspace", "页面如何读取掌握状态、历史和问答记录。"),
        ],
        [3.6, 12.7],
    )

    document.add_heading("10 尚未接入的接口", level=1)
    for text in [
        "source_link 目前是本地样例，正式链接由视频或课堂平台提供。",
        "主持人事件当前用 Chat 命令模拟，未来由平台事件接口触发。",
        "老师端 PDF、网址和接口材料尚未真实接入。",
        "/research 目前只有占位逻辑，尚未接入真实检索。",
        "学生 workspace 页面尚未实现，但数据文件已经准备为共享数据层。",
    ]:
        add_bullet(document, text)

    add_body(
        document,
        "本文的主线是运作逻辑和数据边界。详细分表草案放在 DATABASE-MODEL.md 和 DATABASE-SCHEMA.sql，仅作参考，不要求照搬。",
        after=0,
    )

    document.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(OUTPUT)
